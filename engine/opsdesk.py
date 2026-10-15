"""ops-tickets engine: health checks and alerts in, tickets out.

    python opsdesk.py init                       # write a starter config (~/.ops-tickets/config.json)
    python opsdesk.py status                     # config + which sources/packs are on, as JSON
    python opsdesk.py test-source <database|urls>  # can we reach it? JSON {ok, message}
    python opsdesk.py config-get [dotted.key]
    python opsdesk.py config-set <dotted.key> <json value>
    python opsdesk.py run [--days N]             # run every enabled pack's checks
    python opsdesk.py ingest-alerts <file|-> [--format slack|json] [--days N]   # alerts -> tickets
    python opsdesk.py list [--all] | show <id> | json | report
    python opsdesk.py fix <id> [note] | reopen <id> | ack <id> [note] | unack <id> | owner <id> <name>
    python opsdesk.py step <id> <n> | verify <id> | notifications | owners
    python opsdesk.py export [path] [--status open|acknowledged|fixed|all] [--owner X] [--source Y]

Everything is read-only against the systems it watches. Tickets live in a local
SQLite file (config "tickets_db"); point it at a shared folder to share a list.

Lifecycle
  open          the problem is happening
  acknowledged  known / owned elsewhere: kept, sorted last, back to open only on a spike
  fixed         a person said so; new evidence after the fix reopens it as a regression
  resolved      not seen for RESOLVE_AFTER_DAYS
"""

import importlib.util
import json
import os
import sqlite3
import sys
from datetime import datetime, timedelta, timezone
from pathlib import Path

ENGINE = Path(__file__).resolve().parent
sys.path.insert(0, str(ENGINE))

RESOLVE_AFTER_DAYS = 3
LOOKBACK_HOURS = 30
SEVERITIES = {"high": 0, "medium": 1, "low": 2}
HOME = Path(os.environ.get("OPS_TICKETS_HOME") or Path.home() / ".ops-tickets")
CONFIG_PATH = HOME / "config.json"

DEFAULT_CONFIG = {
    "tickets_db": str(HOME / "tickets.db"),
    "owners": ["Backend", "Mobile app", "Ops", "Me"],
    # who owns tickets from each source; a pack or a person can override per ticket
    "owner_rules": {"Sentry backend": "Backend", "Sentry mobile": "Mobile app", "AWS": "Backend",
                    "Deploys": "Backend", "URLs": "Backend", "SQL rules": "Ops", "default": "Me"},
    "sources": {
        "database": {"enabled": False, "driver": "postgres", "env_file": "",
                     "keys": {"host": "DB_HOST", "port": "DB_PORT", "dbname": "DB_NAME",
                              "user": "DB_USER", "password": "DB_PASSWORD"},
                     "tunnel": ""},
        "slack": {"enabled": False, "channel_id": "", "channel_name": "", "lookback_hours": 26},
        "sentry": {"enabled": False, "organization": "", "projects": []},
        "aws": {"enabled": False, "region": "", "profile": ""},
        "github": {"enabled": False, "repos": []},
        "urls": {"enabled": False, "list": []},
    },
    "packs": [{"name": "sql_rules"}, {"name": "urls"}],
    "sql_rules": [],
    "notify": {"new_high": True, "regression": True, "spike": True},
    "schedule": "08:00",
}


# ------------------------------------------------------------------------ util


def now():
    return datetime.now(timezone.utc)


def iso(dt):
    return dt.astimezone(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")


def out(obj):
    print(json.dumps(obj, default=str))


def deep_merge(base, extra):
    result = dict(base)
    for k, v in (extra or {}).items():
        result[k] = deep_merge(base[k], v) if isinstance(v, dict) and isinstance(base.get(k), dict) else v
    return result


def load_config():
    saved = json.loads(CONFIG_PATH.read_text(encoding="utf-8")) if CONFIG_PATH.exists() else {}
    return deep_merge(DEFAULT_CONFIG, saved)


def save_config(cfg):
    HOME.mkdir(parents=True, exist_ok=True)
    CONFIG_PATH.write_text(json.dumps(cfg, indent=2), encoding="utf-8")


# ---------------------------------------------------------------------- store


def tickets_db(cfg=None):
    cfg = cfg or load_config()
    path = Path(os.path.expanduser(cfg["tickets_db"]))
    path.parent.mkdir(parents=True, exist_ok=True)
    db = sqlite3.connect(path, timeout=15)
    db.row_factory = sqlite3.Row
    db.executescript("""
        create table if not exists tickets (
            id integer primary key autoincrement, fingerprint text unique not null, check_name text not null,
            title text not null, severity text not null, status text not null default 'open', detail text,
            action text, occurrences integer not null default 1, first_seen text not null, last_seen text not null,
            fixed_at text, fixed_note text, reopen_count integer not null default 0, meta text, steps_done text,
            owner text);
        create table if not exists events (
            id integer primary key autoincrement, ticket_id integer not null, at text not null, kind text not null,
            note text, notified integer not null default 0);
        create table if not exists alerts (source text, key text, at text, data text, primary key (source, key, at));
        create table if not exists ingests (at text, alerts integer, source text);
    """)
    # older stores (pre-plugin ops_monitor): add what is missing, rename the alert check
    cols = lambda t: {r[1] for r in db.execute(f"pragma table_info({t})")}  # noqa: E731
    for col in ("meta", "steps_done", "owner"):
        if col not in cols("tickets"):
            db.execute(f"alter table tickets add column {col} text")
    if "notified" not in cols("events"):
        db.execute("alter table events add column notified integer not null default 0")
    if "source" not in cols("ingests"):
        db.execute("alter table ingests add column source text")
    db.execute("update tickets set check_name='alerts' where check_name='slack_alerts'")
    db.commit()
    return db


def log(db, ticket_id, kind, note=None):
    db.execute("insert into events (ticket_id, at, kind, note) values (?,?,?,?)", (ticket_id, iso(now()), kind, note))


def need(db, tid):
    r = db.execute("select * from tickets where id=?", (tid,)).fetchone()
    if not r:
        sys.exit(f"No ticket #{tid}")
    return r


# --------------------------------------------------------------------- packs


class Context:
    """What a pack's check gets: the config, and lazy access to the sources it needs."""

    def __init__(self, cfg):
        self.cfg = cfg
        self._db = None

    def db(self):
        if self._db is None:
            import sources_db
            self._db = sources_db.connect(self.cfg["sources"]["database"])
        return self._db

    def cursor(self):
        return self.db().cursor()


def load_packs(cfg):
    """Built-in packs by name (engine/packs/<name>.py), others by file path."""
    packs = []
    for spec in cfg.get("packs", []):
        path = Path(spec["path"]).expanduser() if spec.get("path") else ENGINE / "packs" / f"{spec['name']}.py"
        mod_name = f"opspack_{path.stem}"
        loaded = importlib.util.spec_from_file_location(mod_name, path)
        mod = importlib.util.module_from_spec(loaded)
        sys.modules[mod_name] = mod
        loaded.loader.exec_module(mod)
        pack = dict(mod.PACK)
        pack.setdefault("name", path.stem)
        pack["_options"] = spec
        packs.append(pack)
    return packs


def pack_ready(pack, cfg):
    return all(cfg["sources"].get(s, {}).get("enabled") for s in pack.get("needs", []))


def default_owner(cfg, packs, check_name, meta):
    for p in packs:
        fn = p.get("default_owner")
        owner = fn(check_name, meta or {}) if fn else None
        if owner:
            return owner
    rules = cfg["owner_rules"]
    return rules.get(source_of({"meta": meta, "check_name": check_name}), rules.get("default"))


def source_of(t):
    m = t.get("meta") or {}
    return m.get("source") or m.get("pack") or t.get("check_name") or "Checks"


def solution_steps(t, packs):
    m = t.get("meta") or {}
    for p in packs:
        fn = p.get("solution_steps")
        steps = fn(t) if fn else None
        if steps:
            return steps
    steps = list(m.get("steps") or [])
    return steps + ["Press Verify now: it re-checks this problem against the latest data."] if steps else \
        [t.get("action") or "Investigate using the detail above.", "Press Verify now when done."]


# ---------------------------------------------------------------- lifecycle


def apply_finding(db, cfg, packs, name, f, t):
    """Open, refresh, or bring back one ticket from one finding (see module doc)."""
    meta = f.get("meta") or {}
    meta_json = json.dumps(meta, default=str)
    owner = f.get("owner") or default_owner(cfg, packs, name, meta)
    row = db.execute("select * from tickets where fingerprint=?", (f["fingerprint"],)).fetchone()
    if row is None:
        status = "acknowledged" if f.get("known") else "open"
        c = db.execute(
            "insert into tickets (fingerprint, check_name, title, severity, detail, action, first_seen, last_seen,"
            " meta, status, owner) values (?,?,?,?,?,?,?,?,?,?,?)",
            (f["fingerprint"], name, f["title"], f["severity"], f.get("detail"), f.get("action"), t, t, meta_json,
             status, owner))
        log(db, c.lastrowid, "opened")
        if f.get("known"):
            log(db, c.lastrowid, "acknowledged", f["known"])
        return
    was = row["status"]
    if was in ("fixed", "resolved") and iso(f["last_event"]) <= (row["fixed_at"] or row["last_seen"]):
        return  # evidence predates the fix
    status = "open"
    if was in ("fixed", "resolved"):
        log(db, row["id"], "reopened", "regression: seen again after " + was)
    elif was == "acknowledged":
        status = "open" if f.get("spike") else "acknowledged"
        if f.get("spike"):
            log(db, row["id"], "spike", "acknowledged issue is spiking: back to open")
    elif f.get("spike") and "SPIKING" not in (row["title"] or ""):
        log(db, row["id"], "spike", "error rate jumped")
    db.execute(
        "update tickets set title=?, severity=?, detail=?, action=?, meta=?, status=?, last_seen=?,"
        " occurrences=occurrences+1, reopen_count=reopen_count+?, fixed_at=case when ?='open' then null else fixed_at end,"
        " owner=coalesce(owner, ?) where id=?",
        (f["title"], f["severity"], f.get("detail"), f.get("action"), meta_json, status, t,
         1 if was in ("fixed", "resolved") else 0, status, owner, row["id"]))


def auto_resolve(db, names, seen, days=RESOLVE_AFTER_DAYS):
    cutoff = iso(now() - timedelta(days=days))
    marks = ",".join("?" * len(names))
    for r in db.execute(f"select * from tickets where status in ('open','acknowledged') and check_name in ({marks})",
                        list(names)).fetchall():
        if r["fingerprint"] not in seen and r["last_seen"] < cutoff:
            db.execute("update tickets set status='resolved' where id=?", (r["id"],))
            log(db, r["id"], "auto-resolved", f"not seen for {days} days")


def cmd_run(days=None):
    cfg = load_config()
    packs = load_packs(cfg)
    db = tickets_db(cfg)
    ctx = Context(cfg)
    since = now() - (timedelta(days=days) if days else timedelta(hours=LOOKBACK_HOURS))
    t = iso(now())
    seen, names, failed, skipped = set(), [], [], []
    for pack in packs:
        if not pack_ready(pack, cfg):
            skipped.append(pack["name"])
            continue
        for name, fn in pack.get("checks", []):
            names.append(name)
            try:
                findings = list(fn(ctx, since))
            except Exception as exc:  # a broken check is itself a ticket
                if ctx._db is not None:
                    ctx._db.rollback()
                failed.append(name)
                findings = [{"fingerprint": f"check-broken:{name}", "title": f"Check '{name}' failed to run",
                             "severity": "low", "detail": f"{type(exc).__name__}: {exc}",
                             "action": "Fix the check or its source connection.", "last_event": now(),
                             "meta": {"source": "Checks"}}]
            for f in findings:
                if (f.get("meta") or {}).get("missing_in") == "Resolved":
                    continue
                seen.add(f["fingerprint"])
                apply_finding(db, cfg, packs, name, f, t)
    if names:
        auto_resolve(db, names, seen)
    db.commit()
    print(report(db))
    if skipped:
        print(f"\nPacks skipped (their source is not enabled): {', '.join(skipped)}")
    if failed:
        print(f"Checks that errored: {', '.join(failed)}")


def cmd_ingest_alerts(path, fmt="slack", days=3):
    """Alerts collected by the mod (Slack) or by Claude (Sentry/AWS/GitHub connectors) -> tickets."""
    import alerts

    cfg = load_config()
    packs = load_packs(cfg)
    db = tickets_db(cfg)
    raw = sys.stdin.read() if path == "-" else Path(path).read_text(encoding="utf-8")
    parsed = alerts.parse(raw) if fmt == "slack" else alerts.from_json(raw)
    for a in parsed:
        db.execute("insert or ignore into alerts values (?,?,?,?)",
                   (a["source"], a["key"], iso(a["at"]), json.dumps(a, default=str)))
    db.execute("insert into ingests values (?,?,?)", (iso(now()), len(parsed), fmt))
    since = iso(now() - timedelta(days=days))
    recent = []
    for (data,) in db.execute("select data from alerts where at >= ?", (since,)):
        a = json.loads(data)
        a["at"] = datetime.fromisoformat(a["at"])
        recent.append(a)
    rules = [r for p in packs for r in p.get("alert_rules", [])]
    t = iso(now())
    seen = set()
    for f in alerts.findings(recent, rules, cfg["owner_rules"]):
        seen.add(f["fingerprint"])
        apply_finding(db, cfg, packs, "alerts", f, t)
    for r in db.execute("select * from tickets where status in ('open','acknowledged') and check_name='alerts'").fetchall():
        if r["fingerprint"] not in seen:
            db.execute("update tickets set status='resolved' where id=?", (r["id"],))
            log(db, r["id"], "auto-resolved", f"no alert for {days} days")
    db.commit()
    print(f"Read {len(parsed)} alert(s); {len(seen)} alert ticket(s) active.")


# -------------------------------------------------------------------- verify


def cmd_verify(tid):
    cfg = load_config()
    packs = load_packs(cfg)
    db = tickets_db(cfg)
    t = dict(need(db, tid))
    meta = json.loads(t["meta"]) if t.get("meta") else {}
    ok, message = None, None
    if meta.get("kind") == "alert":
        keys = meta.get("alert_keys") or [meta.get("alert_key")]
        last = db.execute(f"select max(at) from alerts where key in ({','.join('?' * len(keys))})", keys).fetchone()[0]
        synced = db.execute("select max(at) from ingests").fetchone()[0]
        day_ago = iso(now() - timedelta(hours=24))
        fresh = bool(synced and synced >= day_ago)
        ok = fresh and (last is None or last < day_ago)
        message = (f"No new alert in the last 24h (alerts last collected {synced}). Safe to mark fixed." if ok
                   else f"Last alert was {last}." if fresh
                   else f"Alerts were last collected {synced or 'never'}: run /ops-run first.")
    else:
        ctx = Context(cfg)
        for p in packs:
            fn = p.get("verify")
            res = fn(ctx, t, meta) if fn else None
            if res:
                ok, message = res
                break
        if ok is None:
            check = {n: f for p in packs for n, f in p.get("checks", [])}.get(t["check_name"])
            since = datetime.fromisoformat(t["last_seen"].replace("Z", "+00:00")) + timedelta(seconds=1)
            found = bool(check) and any(f["fingerprint"] == t["fingerprint"] for f in check(ctx, since))
            ok = not found
            message = ("Nothing new since this ticket was last seen. Safe to mark fixed." if ok
                       else "Still happening: new evidence after the ticket was last updated.")
    log(db, tid, "verified" if ok else "verify-failed", message)
    db.commit()
    out({"ok": ok, "message": message})


# -------------------------------------------------------------------- reading


def order_key(r):
    return ({"open": 0, "acknowledged": 1}.get(r["status"], 2), SEVERITIES.get(r["severity"], 9), r["id"])


def fmt(r, long=False):
    s = f"#{r['id']:<4} [{r['severity']:<6}] {r['status']:<12} {r['title']}  -> {r['owner'] or '?'}"
    if r["reopen_count"]:
        s += f"  (REGRESSION x{r['reopen_count']})"
    if long:
        s += f"\n      seen {r['occurrences']}x, first {r['first_seen']}, last {r['last_seen']}"
        s += f"\n      detail: {(r['detail'] or '').replace(chr(10), chr(10) + '              ')}"
        s += f"\n      action: {r['action']}"
    return s


def report(db):
    rs = sorted(db.execute("select * from tickets where status='open'").fetchall(), key=order_key)
    lines = [f"# Ops tickets - {now():%Y-%m-%d %H:%M} UTC", ""]
    lines += [f"{len(rs)} open ticket(s)", ""] + [fmt(r, long=True) + "\n" for r in rs] if rs else ["No open tickets."]
    acked = db.execute("select count(*) from tickets where status='acknowledged'").fetchone()[0]
    fixed = db.execute("select count(*) from tickets where status='fixed'").fetchone()[0]
    lines.append(f"({acked} acknowledged, {fixed} fixed and watched for regressions)")
    return "\n".join(lines)


def tickets_json(cfg, packs, db):
    result = []
    for r in sorted(db.execute("select * from tickets").fetchall(), key=order_key):
        t = dict(r)
        t["meta"] = json.loads(t["meta"]) if t.get("meta") else None
        t["events"] = [dict(e) for e in db.execute(
            "select at, kind, note from events where ticket_id=? order by id", (r["id"],))]
        t["steps"] = solution_steps(t, packs)
        t["steps_done"] = json.loads(t["steps_done"]) if t.get("steps_done") else []
        t["source"] = source_of(t)
        result.append(t)
    return result


def cmd_export(path=None, status="open", owner=None, source=None):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    cfg = load_config()
    packs = load_packs(cfg)
    db = tickets_db(cfg)
    rows = [t for t in tickets_json(cfg, packs, db)
            if (status == "all" and t["status"] != "resolved" or t["status"] == status)
            and (not owner or t["owner"] == owner) and (not source or t["source"] == source)]
    cols = ["Ticket", "Severity", "Status", "Owner", "Source", "Country", "Location", "City", "Problem",
            "Title", "Detail", "Solution steps", "Steps done", "Times seen", "First seen (UTC)", "Last seen (UTC)"]
    wb = Workbook()
    ws = wb.active
    ws.title = "Tickets"
    ws.append(cols)
    fills = {"high": "F8D7DA", "medium": "FFF3CD", "low": "E2E3E5"}
    for t in rows:
        m = t["meta"] or {}
        ws.append([t["id"], t["severity"], t["status"], t["owner"] or "", t["source"], m.get("country", ""),
                   m.get("store", ""), m.get("city", ""), m.get("missing_in", ""), t["title"], t["detail"] or "",
                   "\n".join(f"{'[x]' if i in t['steps_done'] else '[ ]'} {i + 1}. {s}" for i, s in enumerate(t["steps"])),
                   f"{len(t['steps_done'])}/{len(t['steps'])}", t["occurrences"], t["first_seen"], t["last_seen"]])
        ws.cell(ws.max_row, 2).fill = PatternFill("solid", fgColor=fills.get(t["severity"], "FFFFFF"))
    for c in ws[1]:
        c.font = Font(bold=True)
    for i, w in enumerate([8, 10, 13, 18, 16, 13, 20, 12, 28, 50, 60, 80, 10, 10, 20, 20]):
        ws.column_dimensions[chr(65 + i)].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    dest = Path(path) if path else Path.home() / "Downloads" / f"Ops tickets - {now():%Y-%m-%d %H%M}.xlsx"
    dest.parent.mkdir(parents=True, exist_ok=True)
    wb.save(dest)
    print(f"Exported {len(rows)} ticket(s) to {dest}")


# ---------------------------------------------------------------------- setup


def cmd_status():
    cfg = load_config()
    packs = []
    try:
        packs = [{"name": p["name"], "description": p.get("description", ""), "needs": p.get("needs", []),
                  "ready": pack_ready(p, cfg)} for p in load_packs(cfg)]
    except Exception as exc:
        packs = [{"name": "error", "description": str(exc), "needs": [], "ready": False}]
    db = tickets_db(cfg)
    last = {r["source"]: r["at"] for r in db.execute("select source, max(at) as at from ingests group by source")}
    out({"config_path": str(CONFIG_PATH), "configured": CONFIG_PATH.exists(), "tickets_db": cfg["tickets_db"],
         "sources": {k: {"enabled": v.get("enabled", False)} for k, v in cfg["sources"].items()},
         "packs": packs, "owners": cfg["owners"], "last_alert_ingest": last, "schedule": cfg.get("schedule")})


def cmd_test_source(name):
    cfg = load_config()
    src = cfg["sources"].get(name, {})
    try:
        if name == "database":
            import sources_db
            conn = sources_db.connect(src)
            cur = conn.cursor()
            cur.execute("select 1")
            cur.fetchone()
            out({"ok": True, "message": f"Connected ({src.get('driver')})."})
        elif name == "urls":
            sys.path.insert(0, str(ENGINE / "packs"))
            from urls import probe
            results = [probe(u) for u in src.get("list", [])]
            bad = [r for r in results if not r["ok"]]
            out({"ok": not bad and bool(results),
                 "message": f"{len(results) - len(bad)}/{len(results)} URLs healthy" if results else "No URLs configured."})
        else:
            out({"ok": None, "message": "Checked by the mod (connector detection), not by the engine."})
    except Exception as exc:
        hint = f" Is the tunnel running? ({src.get('tunnel')})" if name == "database" and src.get("tunnel") else ""
        out({"ok": False, "message": f"{type(exc).__name__}: {str(exc).strip()[:300]}.{hint}"})


def set_dotted(cfg, key, value):
    node = cfg
    parts = key.split(".")
    for p in parts[:-1]:
        node = node.setdefault(p, {})
    node[parts[-1]] = value


def main(argv):
    if not argv:
        sys.exit(__doc__)
    cmd, args = argv[0], argv[1:]
    opt = lambda n: args[args.index(n) + 1] if n in args else None  # noqa: E731
    if cmd == "init":
        if CONFIG_PATH.exists() and "--force" not in args:
            print(f"Config already exists: {CONFIG_PATH}")
        else:
            save_config(DEFAULT_CONFIG)
            print(f"Wrote {CONFIG_PATH}")
    elif cmd == "status":
        cmd_status()
    elif cmd == "test-source":
        cmd_test_source(args[0])
    elif cmd == "config-get":
        cfg = load_config()
        for p in (args[0].split(".") if args else []):
            cfg = cfg.get(p) if isinstance(cfg, dict) else None
        out(cfg)
    elif cmd == "config-set":
        cfg = load_config()
        set_dotted(cfg, args[0], json.loads(args[1]))
        save_config(cfg)
        print(f"Set {args[0]}.")
    elif cmd == "run":
        cmd_run(int(opt("--days")) if opt("--days") else None)
    elif cmd == "ingest-alerts":
        cmd_ingest_alerts(args[0], opt("--format") or "slack", int(opt("--days") or 3))
    elif cmd == "json":
        cfg = load_config()
        out(tickets_json(cfg, load_packs(cfg), tickets_db(cfg)))
    elif cmd == "list":
        db = tickets_db()
        q = "select * from tickets" + ("" if "--all" in args else " where status in ('open','acknowledged')")
        rs = sorted(db.execute(q).fetchall(), key=order_key)
        print("\n".join(fmt(r) for r in rs) or "No tickets.")
    elif cmd == "show":
        db = tickets_db()
        r = need(db, int(args[0]))
        print(fmt(r, long=True))
        for e in db.execute("select * from events where ticket_id=? order by id", (r["id"],)):
            print(f"      - {e['at']} {e['kind']} {e['note'] or ''}")
    elif cmd == "report":
        print(report(tickets_db()))
    elif cmd in ("fix", "reopen", "ack", "unack"):
        db = tickets_db()
        r = need(db, int(args[0]))
        note = " ".join(args[1:]) or None
        status = {"fix": "fixed", "reopen": "open", "ack": "acknowledged", "unack": "open"}[cmd]
        if cmd == "fix":
            db.execute("update tickets set status='fixed', fixed_at=?, fixed_note=? where id=?", (iso(now()), note, r["id"]))
        else:
            db.execute("update tickets set status=?, fixed_at=case when ?='open' then null else fixed_at end where id=?",
                       (status, status, r["id"]))
        log(db, r["id"], {"fix": "fixed", "reopen": "reopened", "ack": "acknowledged", "unack": "unacknowledged"}[cmd], note)
        db.commit()
        print(f"#{r['id']} is now {status}.")
    elif cmd == "owner":
        db = tickets_db()
        r = need(db, int(args[0]))
        owner = " ".join(args[1:]).strip()
        db.execute("update tickets set owner=? where id=?", (owner, r["id"]))
        log(db, r["id"], "owner", f"{r['owner'] or '-'} -> {owner}")
        db.commit()
        print(f"#{r['id']} owner: {owner}.")
    elif cmd == "owners":
        out(load_config()["owners"])
    elif cmd == "step":
        db = tickets_db()
        r = need(db, int(args[0]))
        n = int(args[1])
        done = set(json.loads(r["steps_done"] or "[]")) ^ {n}
        db.execute("update tickets set steps_done=? where id=?", (json.dumps(sorted(done)), r["id"]))
        log(db, r["id"], "step-done" if n in done else "step-undone", f"step {n + 1}")
        db.commit()
        print(f"#{r['id']} step {n + 1} {'done' if n in done else 'not done'}.")
    elif cmd == "verify":
        cmd_verify(int(args[0]))
    elif cmd == "notifications":
        cfg = load_config()
        db = tickets_db(cfg)
        want = cfg["notify"]
        kinds = [k for k, on in (("reopened", want.get("regression")), ("spike", want.get("spike"))) if on]
        rs = db.execute(f"""
            select e.kind, e.at, t.id as ticket, t.title, t.severity, t.owner from events e join tickets t on t.id = e.ticket_id
            where e.notified = 0 and (e.kind in ({','.join("?" * len(kinds)) or "''"})
                  or (? and e.kind = 'opened' and t.severity = 'high' and t.status = 'open')) order by e.id""",
            [*kinds, 1 if want.get("new_high") else 0]).fetchall()
        db.execute("update events set notified = 1 where notified = 0")
        db.commit()
        out([dict(r) for r in rs])
    elif cmd == "export":
        flags = {"--status", "--owner", "--source"}
        rest = [a for i, a in enumerate(args) if a not in flags and (i == 0 or args[i - 1] not in flags)]
        cmd_export(rest[0] if rest else None, opt("--status") or "open", opt("--owner"), opt("--source"))
    else:
        sys.exit(__doc__)


if __name__ == "__main__":
    main(sys.argv[1:])
